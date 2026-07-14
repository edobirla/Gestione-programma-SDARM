"""Script una-tantum per creare inni di esempio nella cartella 'database di esempio'."""
import os
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

HYMNS = [
    ("001", "Lode a Dio", [
        "Lode a Dio\nnell'alto dei cieli",
        "Strofa 1\nLode a te Signore\nche regni nei cieli\nper sempre con noi",
        "Ritornello\nAlleluia, alleluia\nlode a te o Signore\nalleluia",
        "Strofa 2\nTu sei la nostra luce\nil nostro cammino\nnella notte oscura",
    ]),
    ("002", "Alleluia", [
        "Alleluia",
        "Strofa 1\nAlleluia, canta il cuore\nalleluia, canta l'anima\nalleluia al Signor",
        "Strofa 2\nNel tuo nome ci raccogliamo\nnel tuo spirito camminiamo\nalleluia al Signor",
    ]),
    ("025", "Gioia nel Signore", [
        "Gioia nel Signore",
        "Strofa 1\nLa gioia del Signore è la mia forza\nla sua luce illumina il mio cuore\nnel suo amore cammino ogni giorno",
        "Ritornello\nGioia, gioia nel Signore\ngioia nel suo santo nome\ngioia piena nel mio cuore",
        "Strofa 2\nQuando il buio sembra avvicinarsi\ne la strada sembra troppo lunga\nla sua mano mi sorregge sempre",
    ]),
    ("047", "Santo Santo Santo", [
        "Santo Santo Santo",
        "Santo, Santo, Santo\nè il Signore Dio dell'universo\ni cieli e la terra sono pieni della sua gloria",
        "Osanna nell'alto dei cieli\nbendetto colui che viene\nnel nome del Signore\nosanna nell'alto dei cieli",
    ]),
    ("088", "Anima Christi", [
        "Anima Christi",
        "Anima di Cristo santificami\nCorpo di Cristo salvami\nSangue di Cristo inebriami",
        "Acqua del costato di Cristo lavami\nPassione di Cristo confortami\nO buon Gesù esaudiscimi",
        "Nel tuo profondo mi nascondi\ndal maligno proteggimi\nnell'ora della morte chiamami",
        "E comanda che io venga a te\nacciocché con i tuoi Santi ti lodi\nnei secoli dei secoli. Amen.",
    ]),
    ("103", "Padre Nostro", [
        "Padre Nostro",
        "Padre nostro che sei nei cieli\nsia santificato il tuo nome\nvenga il tuo regno",
        "Sia fatta la tua volontà\ncome in cielo così in terra\ndacci oggi il nostro pane quotidiano",
        "E rimetti a noi i nostri debiti\ncome noi li rimettiamo\nai nostri debitori",
        "E non ci indurre in tentazione\nma liberaci dal male\nAmen",
    ]),
]

folder = os.path.join(os.path.dirname(__file__), "database di esempio")
os.makedirs(folder, exist_ok=True)

for number, title, slides in HYMNS:
    prs = Presentation()
    prs.slide_width = prs.slide_width  # default 10"
    blank_layout = prs.slide_layouts[6]  # blank

    for slide_text in slides:
        slide = prs.slides.add_slide(blank_layout)
        txBox = slide.shapes.add_textbox(
            left=914400, top=457200,
            width=prs.slide_width - 1828800,
            height=prs.slide_height - 914400
        )
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = slide_text
        p.runs[0].font.size = Pt(28)
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    filename = f"{number} - {title}.pptx"
    prs.save(os.path.join(folder, filename))
    print(f"Creato: {filename}")

print("\nDatabase di esempio creato con successo!")
